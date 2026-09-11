"""入力シートのテンプレート(.xlsx)を生成する。

列定義は src/excel/columns.json を唯一の正とし、アプリの取り込み処理と同じ定義から作る。

    pip install openpyxl && python3 scripts/make-template.py
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SPEC = json.loads((ROOT / "src/excel/columns.json").read_text(encoding="utf-8"))
OUT = ROOT / "public" / SPEC["fileName"]

HEADER_FILL = PatternFill("solid", fgColor="E8EEF7")
REQUIRED_FILL = PatternFill("solid", fgColor="FFF3CD")   # 必須列の見出し
EXAMPLE_FILL = PatternFill("solid", fgColor="F4F3F0")    # 記入例の行
THIN = Side(style="thin", color="C4C2BA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FMT = {"yen": "#,##0", "area": "0.00", "minutes": "0", "floor": "0"}


def build() -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet in SPEC["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        cols = sheet["columns"]

        # 1行目: 見出し（必須列に色を付ける）
        for i, col in enumerate(cols, start=1):
            c = ws.cell(row=1, column=i, value=col["header"] + ("  ●" if col["required"] else ""))
            c.font = Font(bold=True, size=10)
            c.fill = REQUIRED_FILL if col["required"] else HEADER_FILL
            c.border = BORDER
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = col["width"]
            if col.get("note"):
                c.comment = None  # コメントは環境差が出るため使わない

        # 2行目: 記入例（消して使う）
        for i, col in enumerate(cols, start=1):
            c = ws.cell(row=2, column=i, value=col["example"])
            c.fill = EXAMPLE_FILL
            c.font = Font(italic=True, size=9, color="7A776E")
            c.border = BORDER
            if col["kind"] in FMT:
                c.number_format = FMT[col["kind"]]

        ws.freeze_panes = f"A{sheet['firstDataRow']}"
        ws.row_dimensions[1].height = 30

        # 説明を右側の余白に置く
        note_col = len(cols) + 2
        ws.cell(row=1, column=note_col, value=sheet["description"]).font = Font(size=9, color="555250")
        ws.cell(row=2, column=note_col, value="● の付いた列は必須です。2行目は記入例なので消してから使ってください。").font = Font(size=9, color="555250")
        ws.cell(row=3, column=note_col, value=SPEC["note"]).font = Font(size=9, color="555250")
        notes = [f'・{c["header"]}: {c["note"]}' for c in cols if c.get("note")]
        for n, text in enumerate(notes):
            ws.cell(row=4 + n, column=note_col, value=text).font = Font(size=9, color="555250")
        ws.column_dimensions[get_column_letter(note_col)].width = 56

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"生成: {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    build()
