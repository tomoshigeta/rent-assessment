"""実際の募集図面6件を記入した動作確認用シートを作る。

数値は麻布十番エリアの実際の募集図面(2026-09)から転記したもの。
対象物件の現在家賃は図面に無いため仮定値で、docs/spec.md に明記してある。

    python3 scripts/make-template.py && python3 scripts/make-sample.py
"""
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "public" / "更新賃料検討_入力シート.xlsx"
OUT = ROOT / "public" / "更新賃料検討_入力シート_記入例_麻布十番.xlsx"

SUBJECT = ["麻布十番ロイヤルプレイス 604号室", 210000, 0, 38.10,
           "東京都港区麻布十番", "東京メトロ南北線・都営大江戸線", "麻布十番", 2, None, 6]
# 築年月は入手できていないため空欄。docs/spec.md 第2章の仮定一覧を参照。
LISTINGS = [
    ["アーバンパーク麻布十番 0902号室", 310000, 10000, 42.14, "東京都港区麻布十番2-12-12", "東京メトロ南北線", "麻布十番", 3, "1999-11", 9],
    ["カスタリア麻布十番 705号室",      280000, 10000, 42.84, "東京都港区麻布十番2-10-1",  "東京メトロ南北線", "麻布十番", 3, "2005-09", 7],
    ["カスタリア麻布十番 603号室",      241000, 10000, 36.95, "東京都港区麻布十番2-10-1",  "東京メトロ南北線", "麻布十番", 3, "2005-09", 6],
    ["南麻布1-5-8 702号室",            185000,     0, 39.00, "東京都港区南麻布1-5-8",     "東京メトロ南北線", "麻布十番", 4, "1970",    7],
    ["メゾン東麻布 401号室",           178000, 12000, 38.00, "東京都港区東麻布2-22-10",   "都営大江戸線",     "赤羽橋",   3, "1983",    4],
    ["イイダアネックス麻布十番 501号室", 170000,  5000, 43.06, "東京都港区東麻布3-7-8",     "東京メトロ南北線", "麻布十番", 2, "1992-08", 5],
]

wb = load_workbook(SRC)
for name, rows in (("対象物件", [SUBJECT]), ("比較事例", LISTINGS)):
    ws = wb[name]
    # 2行目は記入例として残し、3行目から書く（columns.json の firstDataRow に合わせる）
    for r, row in enumerate(rows, start=3):
        for c, v in enumerate(row, start=1):
            # openpyxl の cell(value=None) は既存値を消さないため、必ず代入する
            cell = ws.cell(row=r, column=c)
            cell.value = v
            cell.font = Font(size=10)
wb.save(OUT)
print(f"生成: {OUT.relative_to(ROOT)}")
